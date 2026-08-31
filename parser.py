#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
신제품(NEW-IN) 판매 대시보드 재사용 파서.

로컬 파일만 파싱한다(네트워크 접근 없음). 다운로드/복호화 결과 경로만 받는다.

CLI:
  python parser.py --manifest manifest.json --release release.xlsx \
                   --keys keys.txt --out new_records.json

manifest.json = [{"path": "...", "kind": "cafe24|naver_sales|naver_refund|29cm|musinsa",
                  "channel": "자사몰|네이버|29cm|무신사"(선택)}]

발매정보(release.xlsx)로 11자리 바코드(컬러단위, 뉴컬러만 등록) 매칭.
판매/환불 모두 발매정보 11자리 정확 일치할 때만 반영(상품7은 고유키/네이버 폴백용).
멱등: keys.txt(기존 고유키) 및 내부중복 dedup. 결정적 출력.
"""
import argparse, csv, io, json, sys, datetime
import openpyxl

# --------------------------------------------------------------------------- #
# helpers
# --------------------------------------------------------------------------- #
KIND_CHANNEL = {
    'cafe24': '자사몰',
    'naver_sales': '네이버',
    'naver_refund': '네이버',
    '29cm': '29cm',
    'musinsa': '무신사',
}

def toint(v):
    """콤마/소수점 포함 금액 문자열을 반올림 정수로."""
    if v is None or v == '':
        return 0
    try:
        return int(round(float(str(v).replace(',', '').strip())))
    except Exception:
        return 0

def brand_of(code):
    c = str(code)[:1].upper()
    if c == 'R':
        return 'RAWROW'
    if c in ('N', 'J'):
        return 'NAUTICA'
    return ''

def norm_date_from_num(s):
    """YYYYMMDD 앞 8자리 -> YYYY-MM-DD."""
    s = str(s)
    return '%s-%s-%s' % (s[0:4], s[4:6], s[6:8])

def parse_date_cell(v):
    """날짜/날짜문자 셀 -> YYYY-MM-DD (없으면 '')."""
    if v is None or v == '':
        return ''
    if isinstance(v, datetime.datetime):
        return v.date().strftime('%Y-%m-%d')
    if isinstance(v, datetime.date):
        return v.strftime('%Y-%m-%d')
    s = str(v).strip().replace('.', '-').replace('/', '-')
    return s[:10]

def to_date_obj(v):
    if isinstance(v, datetime.datetime):
        return v.date()
    if isinstance(v, datetime.date):
        return v
    try:
        s = str(v).strip().replace('.', '-').replace('/', '-')
        return datetime.datetime.strptime(s[:10], '%Y-%m-%d').date()
    except Exception:
        return None

def load_xlsx(path):
    """암호화 xlsx면 2222 먼저·실패시 3333으로 복호화. 아니면 그대로."""
    import msoffcrypto
    with open(path, 'rb') as f:
        head = f.read(8)
    # CFB/OLE header => 암호화(CDFV2) 가능성
    if head[:4] == b'\xd0\xcf\x11\xe0':
        last = None
        for pw in ('2222', '3333'):
            try:
                with open(path, 'rb') as f:
                    off = msoffcrypto.OfficeFile(f)
                    off.load_key(password=pw)
                    buf = io.BytesIO()
                    off.decrypt(buf)
                    buf.seek(0)
                    return openpyxl.load_workbook(buf, data_only=True)
            except Exception as e:
                last = e
        raise RuntimeError('복호화 실패(2222/3333): %s' % last)
    return openpyxl.load_workbook(path, data_only=True)

def pick_sheet(wb, preferred):
    for nm in preferred:
        if nm in wb.sheetnames:
            return wb[nm]
    return wb[wb.sheetnames[0]]

def header_index(header_row, name):
    """헤더에서 이름(공백/좌우 trim 무시) 인덱스. 없으면 None."""
    target = str(name).replace(' ', '').strip()
    for i, h in enumerate(header_row):
        if h is None:
            continue
        if str(h).replace(' ', '').strip() == target:
            return i
    return None

# --------------------------------------------------------------------------- #
# release (발매정보)
# --------------------------------------------------------------------------- #
def load_release(path):
    """11자리 바코드->info, 상품7->info(가장 이른 발매일) 두 룩업 구성."""
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))
    # 컬럼: 0 바코드, 1 시리즈, 2 상품명, 3 컬러, 4 발매일
    by11 = {}
    by7 = {}
    for r in rows[1:]:
        if not r or r[0] is None:
            continue
        bc = str(r[0]).strip()
        if len(bc) != 11:
            continue
        rel = to_date_obj(r[4]) if len(r) > 4 else None
        info = {
            'barcode': bc,
            'series': (r[1] or '') if len(r) > 1 else '',
            'name': (r[2] or '') if len(r) > 2 else '',
            'color': (r[3] or '') if len(r) > 3 else '',
            'rel': rel,
            'relmonth': rel.strftime('%Y-%m') if rel else '',
            'reldate': rel.strftime('%Y-%m-%d') if rel else '',
            'brand': brand_of(bc),
            'p7': bc[:7],
        }
        by11[bc] = info
        p7 = bc[:7]
        if p7 not in by7 or (rel and by7[p7]['rel'] and rel < by7[p7]['rel']):
            by7[p7] = info
    return by11, by7

# --------------------------------------------------------------------------- #
# record emitter (dedup + field order)
# --------------------------------------------------------------------------- #
FIELD_ORDER = ['고유키', '구분', '채널', '날짜', '주문번호', '바코드', '상품7',
               '브랜드', '발매월', '시리즈', '상품명', '컬러', '수량',
               '매출', '환불금액', '발매일']

class Collector:
    def __init__(self, existing_keys):
        self.seen = set(existing_keys)
        self.records = []

    def add(self, kind, channel, order, p7, info, qty, sale, refund, date_str, barcode=''):
        """kind: '판매'/'환불'. info: 발매정보 dict(없으면 None)."""
        brand = info['brand'] if info else ''
        relmonth = info['relmonth'] if info else ''
        series = info['series'] if info else ''
        name = info['name'] if info else ''
        color = info['color'] if info else ''
        reldate = info['reldate'] if info else ''
        if kind == '환불':
            barcode_out, color_out, qty_out, sale_out = '', '', 1, 0
            refund_out = refund
            amt = refund_out
        else:
            barcode_out, color_out, qty_out, sale_out = barcode, color, qty, sale
            refund_out = 0
            amt = sale_out
        key = '%s|%s|%s|%s|%s|%s|%s' % (kind, channel, order, p7, date_str, qty_out, amt)
        if key in self.seen:
            return False
        self.seen.add(key)
        rec = {
            '고유키': key, '구분': kind, '채널': channel, '날짜': date_str,
            '주문번호': order, '바코드': barcode_out, '상품7': p7, '브랜드': brand,
            '발매월': relmonth, '시리즈': series, '상품명': name, '컬러': color_out,
            '수량': qty_out, '매출': sale_out, '환불금액': refund_out, '발매일': reldate,
        }
        self.records.append(rec)
        return True

# --------------------------------------------------------------------------- #
# parsers per kind
# --------------------------------------------------------------------------- #
def parse_cafe24(path, channel, by11, col, report):
    raw = open(path, 'rb').read().decode('utf-8-sig')
    rd = list(csv.reader(io.StringIO(raw)))
    if not rd:
        return
    hdr = rd[0]
    i_ord = header_index(hdr, '주문번호')
    i_item = header_index(hdr, '자체품목코드')
    i_qty = header_index(hdr, '수량')
    i_cxl = header_index(hdr, '취소구분')
    i_pay = header_index(hdr, '총 결제금액')
    i_ref = header_index(hdr, '총 실제 환불금액')
    if None in (i_ord, i_item, i_cxl, i_pay, i_ref):
        report.setdefault('warn', []).append('cafe24 헤더 컬럼 누락: %s' % path)
        return
    for r in rd[1:]:
        if len(r) <= max(i_ord, i_item, i_qty or 0, i_cxl, i_pay, i_ref):
            continue
        ordnum = str(r[i_ord]).strip()
        if not ordnum:
            continue
        item = str(r[i_item]).strip()
        if '+' in item:            # '...+사은품' -> '+' 앞 11자리
            item = item.split('+')[0]
        bc = item[:11]
        cxl = str(r[i_cxl]).strip()
        qty = toint(r[i_qty]) if i_qty is not None else 1
        pay = toint(r[i_pay])
        refund = toint(r[i_ref])
        date = norm_date_from_num(ordnum[:8])   # 주문일시 #### 대비 주문번호로 날짜 복구
        if len(bc) != 11 or bc not in by11:
            continue
        info = by11[bc]
        p7 = bc[:7]
        if cxl == '취소안함':
            col.add('판매', channel, ordnum, p7, info, qty, pay, 0, date, barcode=bc)
        elif cxl in ('취소', '부분취소'):
            col.add('환불', channel, ordnum, p7, info, 1, 0, refund, date)

def parse_naver_sales(path, channel, by11, col, L, report):
    """판매 레코드 생성 + L(상품주문번호->옵션관리코드11) 구성. 헤더=2번째 행."""
    wb = load_xlsx(path)
    ws = pick_sheet(wb, ['발주발송관리'])
    rows = list(ws.iter_rows(values_only=True))
    if len(rows) < 3:
        return
    hdr = rows[1]
    i_ord = header_index(hdr, '상품주문번호')
    i_opt = header_index(hdr, '옵션관리코드')
    i_amt = header_index(hdr, '최종 상품별 총 주문금액')
    i_qty = header_index(hdr, '수량')
    if None in (i_ord, i_opt, i_amt):
        report.setdefault('warn', []).append('naver_sales 헤더 컬럼 누락: %s' % path)
        return
    for r in rows[2:]:
        if i_ord >= len(r) or r[i_ord] is None:
            continue
        ordnum = str(r[i_ord]).strip()
        opt = r[i_opt] if i_opt < len(r) else None
        opt = str(opt).strip() if opt is not None else ''
        if len(opt) == 11:
            L[ordnum] = opt                        # L: 전체 판매에서 구성
        amt = toint(r[i_amt]) if i_amt < len(r) else 0
        qty = (toint(r[i_qty]) if i_qty is not None and i_qty < len(r) else 0) or 1
        date = norm_date_from_num(ordnum[:8])
        if len(opt) == 11 and opt in by11:
            col.add('판매', channel, ordnum, opt[:7], by11[opt], qty, amt, 0, date, barcode=opt)

def parse_naver_refund(path, channel, by11, by7, col, L, report):
    """환불. 반품 처리상태=='반품완료'만. 11자리 없음 -> L 조인, 실패시 판매자상품코드[:7] 폴백."""
    wb = load_xlsx(path)
    ws = pick_sheet(wb, ['반품관리'])
    rows = list(ws.iter_rows(values_only=True))
    if len(rows) < 2:
        return
    hdr = rows[0]
    i_ord = header_index(hdr, '상품주문번호')
    i_stat = header_index(hdr, '반품 처리상태')
    i_ref = header_index(hdr, '환불일')
    i_scode = header_index(hdr, '판매자 상품코드')
    if None in (i_ord, i_stat, i_ref):
        report.setdefault('warn', []).append('naver_refund 헤더 컬럼 누락: %s' % path)
        return
    for r in rows[1:]:
        if i_ord >= len(r) or r[i_ord] is None:
            continue
        if str(r[i_stat]).strip() != '반품완료':
            continue
        ordnum = str(r[i_ord]).strip()
        refdate = to_date_obj(r[i_ref]) if i_ref < len(r) else None
        date_str = refdate.strftime('%Y-%m-%d') if refdate else ''
        bc = L.get(ordnum)
        if bc and bc in by11:
            col.add('환불', channel, ordnum, bc[:7], by11[bc], 1, 0, 0, date_str)
            continue
        # 폴백: 판매자 상품코드 앞7 == 발매정보 상품7 AND 환불일 >= 해당 상품7 발매일
        scode = r[i_scode] if (i_scode is not None and i_scode < len(r)) else None
        p7 = str(scode).strip()[:7] if scode else ''
        if p7 and p7 in by7 and refdate and by7[p7]['rel'] and refdate >= by7[p7]['rel']:
            col.add('환불', channel, ordnum, p7, by7[p7], 1, 0, 0, date_str)

def parse_29cm(path, channel, by11, col, report):
    """29CM. 판매/반품 파일 공통(단일 시트, 헤더 이름 기반). 컬럼 위치 변동에 견고."""
    import os
    wb = load_xlsx(path)
    ws = pick_sheet(wb, ['주문목록'])          # 없으면 첫 시트(Sheet1)
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return
    hidx = 0
    for k in range(min(5, len(rows))):
        if header_index(rows[k], '주문번호') is not None:
            hidx = k; break
    hdr = rows[hidx]
    i_ord  = header_index(hdr, '주문번호')
    i_stat = header_index(hdr, '주문상태')
    i_bc   = header_index(hdr, '풀바코드')
    i_qty  = header_index(hdr, '수량')
    i_sale = header_index(hdr, '실판매액')       # '실 판매액'(공백무시) 매칭
    if i_sale is None:
        i_sale = header_index(hdr, '판매액')
    i_date = header_index(hdr, '주문일시')
    if None in (i_ord, i_stat, i_bc):
        report.setdefault('warn', []).append('29cm 헤더 컬럼 누락: %s' % path)
        return
    is_refund_file = '반품' in os.path.basename(path)
    def cell(r, i):
        return r[i] if (i is not None and i < len(r)) else None
    for r in rows[hidx + 1:]:
        if not r or cell(r, i_ord) is None:
            continue
        ordnum = str(cell(r, i_ord)).strip()
        if not ordnum:
            continue
        stat = str(cell(r, i_stat) or '').strip()
        bc = str(cell(r, i_bc) or '').strip()
        date = parse_date_cell(cell(r, i_date))
        if not date:
            digits = ''.join(ch for ch in ordnum if ch.isdigit())
            if len(digits) >= 8:
                date = norm_date_from_num(digits[:8])
        match = (len(bc) == 11 and bc in by11)
        if not match:
            continue
        if '교환' in stat:
            continue
        if ('취소' in stat) or ('반품' in stat) or ('환불' in stat):
            col.add('환불', channel, ordnum, bc[:7], by11[bc], 1, 0, 0, date)
        else:
            if is_refund_file:
                continue                          # 반품 파일의 정상건은 판매로 넣지 않음
            qty = toint(cell(r, i_qty))
            sale = toint(cell(r, i_sale))
            if qty > 0 and sale > 0:
                col.add('판매', channel, ordnum, bc[:7], by11[bc], qty, sale, 0, date, barcode=bc)

def parse_musinsa(path, channel, by11, col, report):
    # 주문번호 정밀도 유지를 위해 문자열로 취급
    wb = load_xlsx(path)
    ws = pick_sheet(wb, ['주문내역_통합'])
    rows = list(ws.iter_rows(values_only=True))
    # 컬럼: 0 주문일시,1 주문번호,3 주문상태,4 클레임상태,8 상품명,11 풀바코드,13 수량,19 실결제금액
    for r in rows[1:]:
        if not r or len(r) < 2 or r[1] is None:
            continue
        date = parse_date_cell(r[0])
        ordnum = str(r[1]).strip()
        ostat = str(r[3]).strip() if len(r) > 3 and r[3] else ''
        cstat = str(r[4]).strip() if len(r) > 4 and r[4] else ''
        bc = str(r[11]).strip() if len(r) > 11 and r[11] else ''
        qty = toint(r[13]) if len(r) > 13 else 0
        pay = toint(r[19]) if len(r) > 19 else 0
        match = (len(bc) == 11 and bc in by11)
        is_cancel = (ostat == '주문취소' or cstat == '주문취소')
        is_return = (cstat == '환불완료')
        if is_cancel:
            if match:
                col.add('환불', channel, ordnum, bc[:7], by11[bc], 1, 0, 0, date)
        elif is_return:
            if match:
                col.add('환불', channel, ordnum, bc[:7], by11[bc], 1, 0, 0, date)
                if qty > 0 and pay > 0:
                    col.add('판매', channel, ordnum, bc[:7], by11[bc], qty, pay, 0, date, barcode=bc)
        else:
            if ostat != '결제오류' and match and qty > 0 and pay > 0:
                col.add('판매', channel, ordnum, bc[:7], by11[bc], qty, pay, 0, date, barcode=bc)

# --------------------------------------------------------------------------- #
# main
# --------------------------------------------------------------------------- #
def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--manifest', required=True)
    ap.add_argument('--release', required=True)
    ap.add_argument('--keys', required=True)
    ap.add_argument('--out', required=True)
    args = ap.parse_args()

    by11, by7 = load_release(args.release)

    try:
        existing = set(l.strip() for l in open(args.keys, encoding='utf-8') if l.strip())
    except FileNotFoundError:
        existing = set()

    manifest = json.load(open(args.manifest, encoding='utf-8'))
    order = ('cafe24', 'naver_sales', 'naver_refund', '29cm', 'musinsa')
    groups = {k: [] for k in order}
    for m in manifest:
        kind = m.get('kind')
        if kind not in groups:
            continue
        m = dict(m)
        m['channel'] = m.get('channel') or KIND_CHANNEL.get(kind, '')
        groups[kind].append(m)

    col = Collector(existing)
    report = {'skipped': [], 'warn': []}
    L = {}

    # 네이버 판매 전체에서 L(상품주문번호->옵션관리코드11)을 먼저 구성해야 환불 조인이 정확.
    # 그러나 판매 레코드도 이 시점에 함께 생성(결정적 순서: cafe24->naver_sales->naver_refund->29cm->musinsa).
    naver_failed = False

    # 1) cafe24
    for m in groups['cafe24']:
        try:
            parse_cafe24(m['path'], m['channel'], by11, col, report)
        except Exception as e:
            report['warn'].append('cafe24 실패(%s): %s' % (m['path'], e))

    # 2) naver_sales (+ L 구성)
    for m in groups['naver_sales']:
        try:
            parse_naver_sales(m['path'], m['channel'], by11, col, L, report)
        except Exception as e:
            naver_failed = True
            report['skipped'].append('naver_sales skip(%s): %s' % (m['path'], e))

    # 3) naver_refund (L 조인 + 상품7 폴백)
    for m in groups['naver_refund']:
        try:
            parse_naver_refund(m['path'], m['channel'], by11, by7, col, L, report)
        except Exception as e:
            naver_failed = True
            report['skipped'].append('naver_refund skip(%s): %s' % (m['path'], e))

    # 4) 29cm
    for m in groups['29cm']:
        try:
            parse_29cm(m['path'], m['channel'], by11, col, report)
        except Exception as e:
            report['warn'].append('29cm 실패(%s): %s' % (m['path'], e))

    # 5) musinsa
    for m in groups['musinsa']:
        try:
            parse_musinsa(m['path'], m['channel'], by11, col, report)
        except Exception as e:
            report['warn'].append('musinsa 실패(%s): %s' % (m['path'], e))

    if naver_failed:
        sys.stderr.write('[WARN] 네이버 파일 처리 실패(복호화/파싱). 네이버 일부/전체 skip됨.\n')

    json.dump(col.records, open(args.out, 'w', encoding='utf-8'), ensure_ascii=False)

    # 요약 리포트(stderr)
    sale = [r for r in col.records if r['구분'] == '판매']
    ref = [r for r in col.records if r['구분'] == '환불']
    dates = [r['날짜'] for r in col.records if r['날짜']]
    ch = {}
    for r in col.records:
        ch['%s %s' % (r['채널'], r['구분'])] = ch.get('%s %s' % (r['채널'], r['구분']), 0) + 1
    sys.stderr.write('신규 %d건 (판매 %d / 환불 %d)\n' % (len(col.records), len(sale), len(ref)))
    sys.stderr.write('판매 매출합: %d\n' % sum(r['매출'] for r in sale))
    if dates:
        sys.stderr.write('날짜범위: %s ~ %s\n' % (min(dates), max(dates)))
    sys.stderr.write('채널별: %s\n' % json.dumps(ch, ensure_ascii=False))
    for w in report['warn']:
        sys.stderr.write('[warn] %s\n' % w)
    for s in report['skipped']:
        sys.stderr.write('[skip] %s\n' % s)


if __name__ == '__main__':
    main()
